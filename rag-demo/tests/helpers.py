from io import BytesIO

import docx
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter


def make_pdf(pages: list[str]) -> bytes:
    """ASCII-only: встроенный шрифт fpdf не поддерживает кириллицу."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.cell(text=text)
    return bytes(pdf.output())


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def encrypt_pdf(data: bytes, user_password: str) -> bytes:
    """RC4: не требует пакета cryptography."""
    writer = PdfWriter(clone_from=PdfReader(BytesIO(data)))
    writer.encrypt(user_password=user_password, owner_password="owner", algorithm="RC4-128")
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()
